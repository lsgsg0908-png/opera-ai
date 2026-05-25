"""스킬: 엑셀 데이터 가공"""
from core.executor import file_write
import json, csv, io

SKILL_INFO = {"id": "skill_excel_data", "name": "엑셀 데이터 가공", "description": "openpyxl로 엑셀 데이터 분석/가공", "category": "office"}

def register():
    return SKILL_INFO

def execute(action="create", path=None, data=None, **kwargs):
    try:
        from openpyxl import Workbook, load_workbook
        if action == "create":
            wb = Workbook()
            ws = wb.active
            if data:
                for row_idx, row in enumerate(data, 1):
                    for col_idx, val in enumerate(row, 1):
                        ws.cell(row=row_idx, column=col_idx, value=val)
            if path:
                wb.save(path)
                return {"status": "ok", "path": path}
            return {"error": "저장 경로가 필요합니다"}
        elif action == "read" and path:
            wb = load_workbook(path, data_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([str(v) if v is not None else "" for v in row])
            return {"status": "ok", "rows": rows[:100], "total": len(rows)}
        return {"error": f"지원하지 않는 동작: {action}"}
    except ImportError:
        return {"error": "openpyxl이 설치되지 않았습니다"}
    except Exception as e:
        return {"error": str(e)}
